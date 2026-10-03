"""
AyurHerb - PowerPoint Presentation Generator
Generates a 12-slide professional presentation (AyurHerb_Project_Presentation.pptx)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)  # 16:9 Widescreen

    # Color Palette
    PRIMARY_DARK = RGBColor(27, 67, 50)      # Deep Forest Green #1B4332
    ACCENT_GREEN = RGBColor(45, 106, 79)     # Herbal Green #2D6A4F
    LIGHT_GREEN = RGBColor(82, 183, 136)     # Mint Green #52B788
    BG_LIGHT = RGBColor(245, 247, 245)       # Off White #F5F7F5
    TEXT_DARK = RGBColor(33, 37, 41)         # Dark Charcoal #212529
    TEXT_MUTED = RGBColor(108, 117, 125)     # Muted Gray #6C757D
    CARD_BG = RGBColor(255, 255, 255)        # Pure White
    AMBER = RGBColor(217, 119, 6)            # Amber/Gold

    blank_layout = prs.slide_layouts[6]

    def set_slide_bg(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, subtitle_text="Data Science & Soft Computing Lab Mini-Project"):
        # Header banner
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_sub = tf.paragraphs[0]
        p_sub.text = subtitle_text.upper()
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = LIGHT_GREEN

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = PRIMARY_DARK

    def add_card(slide, left, top, width, height, title="", bg_color=CARD_BG, border_color=LIGHT_GREEN):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        return shape

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1, PRIMARY_DARK)

    # Accent decorative box
    add_card(s1, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5), bg_color=PRIMARY_DARK, border_color=LIGHT_GREEN)

    tb1 = s1.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "AYURHERB"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    p2 = tf1.add_paragraph()
    p2.text = "Data-Driven Ayurvedic Wellness Recommendation & Botanical Vision System"
    p2.font.size = Pt(20)
    p2.font.color.rgb = LIGHT_GREEN
    p2.space_before = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "A Multi-Paradigm Soft Computing & Machine Learning Web Application"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(220, 220, 220)
    p3.space_before = Pt(16)

    p4 = tf1.add_paragraph()
    p4.text = "\nStudent: Mayuri Haram  |  Data Science using Python Lab  |  GitHub: Mayuri051/AyurHerb"
    p4.font.size = Pt(13)
    p4.font.bold = True
    p4.font.color.rgb = RGBColor(255, 255, 255)
    p4.space_before = Pt(30)

    # =========================================================================
    # SLIDE 2: Problem Statement & Motivation
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2, BG_LIGHT)
    add_header(s2, "Problem Statement & Project Objectives")

    # Left Card - The Challenge
    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s2.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "The Challenge in Traditional Healthcare"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    bullets = [
        "Ayurvedic knowledge spans thousands of classical texts but lacks structured, accessible digital retrieval systems.",
        "Patients experience complex, free-text symptoms that require intelligent natural language mapping.",
        "Diagnostic uncertainty is inherent in healthcare; rigid rule-based systems fail to model probability and fuzzy dosage.",
        "Non-experts struggle to identify medicinal plants and cannot determine safe home preparation or contraindications."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(10)

    # Right Card - The AyurHerb Solution
    add_card(s2, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s2.shapes.add_textbox(Inches(7.2), Inches(1.9), Inches(5.0), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "The AyurHerb Solution"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    sol_bullets = [
        "NLP Symptom Matcher: Converts free-text patient descriptions into TF-IDF vectors to match 446 authenticated records.",
        "Ensemble ML & Probability: Combines Random Forest, AdaBoost, and Bayesian Networks for diagnostic validation.",
        "Mamdani Fuzzy Logic: Calculates personalized herbal dosage potency based on symptom severity, duration, and digestion.",
        "Botanical Leaf Vision: Instantly classifies medicinal plants from photos with Sanskrit names and safe brewing recipes."
    ]
    for b in sol_bullets:
        p = tf.add_paragraph()
        p.text = "✔ " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 3: Dataset & Preprocessing
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3, BG_LIGHT)
    add_header(s3, "Dataset Architecture & Preprocessing Pipeline")

    # 4 Stat Cards at the Top
    stats = [
        ("446", "Clinical Records", Inches(0.8)),
        ("34", "Clinical Attributes", Inches(3.8)),
        ("367", "Disease Classes", Inches(6.8)),
        ("120+", "Herbal Formulations", Inches(9.8))
    ]
    for val, label, left in stats:
        add_card(s3, left, Inches(1.6), Inches(2.7), Inches(1.3), border_color=ACCENT_GREEN)
        tb = s3.shapes.add_textbox(left + Inches(0.1), Inches(1.7), Inches(2.5), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = val
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_DARK
        p.alignment = PP_ALIGN.CENTER
        p2 = tf.add_paragraph()
        p2.text = label
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.alignment = PP_ALIGN.CENTER

    # Lower Card: Preprocessing Details
    add_card(s3, Inches(0.8), Inches(3.2), Inches(11.7), Inches(3.6))
    tb = s3.shapes.add_textbox(Inches(1.1), Inches(3.4), Inches(11.1), Inches(3.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Data Cleaning & Composite Feature Engineering"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    pts = [
        "Dataset Source: Curated AyurGenixAI Ayurvedic clinical corpus mapping diseases to herbs, doshas, and therapies.",
        "Unified_Search_Text Generation: Concatenates Disease + Symptoms + Ayurvedic Herbs + Formulation + Doshas for rich n-gram text modeling.",
        "Multilingual Vernacular Support: Normalizes English medical terms with Hindi (हिंदी) and Marathi (मराठी) vernacular nomenclature.",
        "Clinical Red-Flag Annotations: Pre-tags acute emergency conditions (cardiac arrest, respiratory failure, severe trauma) for immediate safety bypass."
    ]
    for pt in pts:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 4: System Architecture Flow
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4, BG_LIGHT)
    add_header(s4, "End-to-End System Architecture")

    stages = [
        ("1. Input & Safety", "• Free-text symptoms\n• Emergency Triage filter\n• Red-flag intercept", Inches(0.8)),
        ("2. NLP Vectorizer", "• TF-IDF n-grams (1,2)\n• 2,897 vocabulary nodes\n• Cosine Similarity (>=0.12)", Inches(3.2)),
        ("3. ML & Bayesian", "• Random Forest (120 trees)\n• AdaBoost ensemble\n• Bayesian Posterior CPT", Inches(5.6)),
        ("4. Fuzzy Engine", "• Severity, Duration, Agni\n• Mamdani Centroid\n• Potency & Vehicle", Inches(8.0)),
        ("5. Patient UI", "• Disease & % Match\n• Herbs & Preparation\n• Diet, Yoga, Prevention", Inches(10.4))
    ]
    for title, desc, left in stages:
        add_card(s4, left, Inches(1.8), Inches(2.15), Inches(4.8), border_color=ACCENT_GREEN)
        tb = s4.shapes.add_textbox(left + Inches(0.15), Inches(2.0), Inches(1.85), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_DARK
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_DARK
        p2.space_before = Pt(12)

    # =========================================================================
    # SLIDE 5: TF-IDF & Cosine Similarity Engine
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5, BG_LIGHT)
    add_header(s5, "Algorithm 1: TF-IDF & Cosine Similarity Engine")

    add_card(s5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s5.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Mathematical Formulation"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    bullets = [
        "TF-IDF converts query and clinical documents into numerical vectors weighted by term frequency and inverse document frequency.",
        "Sublinear TF Scaling: Replaces raw tf with 1 + log(tf) to dampen the impact of repeated words.",
        "Cosine Similarity: Measures the cosine of the angle between query vector q and document vector d in high-dimensional space:",
        "    Cosine(q, d) = (q · d) / (||q|| × ||d||)",
        "Deterministic, real-time matching with zero external API calls or cloud latency."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(8)

    # Right Card: Hyperparameters Table
    add_card(s5, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s5.shapes.add_textbox(Inches(7.2), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Engine Hyperparameters & Configuration"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    params = [
        ("Module File", "recommendation_engine.py"),
        ("ngram_range", "(1, 2) [Unigrams & Bigrams]"),
        ("sublinear_tf", "True (Logarithmic term weighting)"),
        ("max_features", "2,897 distinct n-gram tokens"),
        ("Indexed Docs", "446 authentic clinical records"),
        ("Threshold", "0.12 (Filters irrelevant queries)"),
        ("Execution Time", "< 15 ms per search query")
    ]
    for k, v in params:
        p = tf.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 6: Random Forest & AdaBoost Supervised ML
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6, BG_LIGHT)
    add_header(s6, "Algorithm 2: Random Forest & AdaBoost Ensemble")

    # Left: Random Forest
    add_card(s6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s6.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Random Forest Classifier (Bagging)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    rf_pts = [
        "Builds 120 de-correlated decision trees trained on bootstrap samples with random feature subspaces.",
        "n_estimators: 120 decision trees",
        "max_depth: 20 (Prevents tree overfitting)",
        "class_weight: 'balanced' (Compensates for multi-class sparsity across 367 disease labels)",
        "Training Accuracy: ~93.50%",
        "Role in App: Acts as the primary diagnostic voting model for multi-class symptom prediction."
    ]
    for pt in rf_pts:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(8)

    # Right: AdaBoost
    add_card(s6, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s6.shapes.add_textbox(Inches(7.2), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "AdaBoost Classifier (Sequential Boosting)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    ada_pts = [
        "Trains weak decision tree learners sequentially, adaptively increasing weights of misclassified samples.",
        "n_estimators: 80 boosting rounds",
        "base_estimator: DecisionTree(max_depth=3)",
        "learning_rate: 0.8 (Shrinkage factor)",
        "Algorithm: SAMME (Multi-class boosting)",
        "Role in App: Provides independent cross-verification to validate Random Forest predictions."
    ]
    for pt in ada_pts:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 7: Bayesian Network Uncertainty Inference
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7, BG_LIGHT)
    add_header(s7, "Algorithm 3: Bayesian Network (Uncertainty Modeling)")

    add_card(s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s7.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Probabilistic Medical Belief Network"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    bayes_pts = [
        "Medical diagnosis is inherently uncertain; symptoms can indicate multiple conditions with varying likelihoods.",
        "Bayes' Theorem Formulation:",
        "    P(D | S1..Sn) ∝ P(D) × ∏ P(Si | D)",
        "Prior Probability P(D): Calculated from relative frequency of diseases in the clinical dataset.",
        "Conditional Probability Tables (CPTs): Formed across 367 disease nodes and 462 symptom tokens.",
        "Laplace Smoothing (α = 0.1): Ensures zero-probability protection for rare/unseen symptom-disease combinations."
    ]
    for pt in bayes_pts:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(7)

    # Right: Parameters
    add_card(s7, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s7.shapes.add_textbox(Inches(7.2), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Bayesian Engine Specifications"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    spec = [
        ("Module File", "bayesian_engine.py"),
        ("Disease Nodes", "367 conditioned targets"),
        ("Symptom Nodes", "462 unique symptom vocabulary tokens"),
        ("Smoothing Alpha", "0.1 (Laplace pseudo-counts)"),
        ("Inference Method", "Exact Joint Posterior Computation"),
        ("Dosha Conditioning", "Incorporates Vata/Pitta/Kapha evidence"),
        ("Output in App", "Estimated Diagnostic Likelihood %")
    ]
    for k, v in spec:
        p = tf.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 8: Mamdani Fuzzy Logic Dosage System
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8, BG_LIGHT)
    add_header(s8, "Algorithm 4: Mamdani Fuzzy Logic Dosage Controller")

    add_card(s8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s8.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Fuzzy Membership & Linguistic Sets"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    fuzzy_pts = [
        "Translates imprecise human inputs ('mild pain', 'several days') into continuous membership degrees [0, 1].",
        "Input 1: Severity Score [0–10] (Mild, Moderate, Severe)",
        "Input 2: Duration [1–60 days] (Acute, Subacute, Chronic)",
        "Input 3: Agni / Digestive Fire [0–10] (Manda, Sama, Tikshna)",
        "Output: Herbal Potency Index [0–100%]",
        "    - Mridu (Gentle Restorative < 35%)",
        "    - Madhyama (Standard Decoction 35–68%)",
        "    - Tikshna (Intensive Concentrate > 68%)"
    ]
    for pt in fuzzy_pts:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(6)

    # Right: Rules & Defuzzification
    add_card(s8, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s8.shapes.add_textbox(Inches(7.2), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Rules & Defuzzification Process"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    f_spec = [
        "Module File: fuzzy_engine.py",
        "Defuzzification: Centroid (Center of Gravity) method",
        "Sample Rule 1: IF Severity is High AND Duration is Chronic THEN Potency is Tikshna",
        "Sample Rule 2: IF Severity is Low AND Agni is Manda THEN Potency is Mridu",
        "Anupana (Vehicle) Engine: Recommends carriers like Honey (Kapha), Ghee (Pitta), or Warm Water (Vata) based on dosha context."
    ]
    for pt in f_spec:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 9: Botanical Vision & Plant Scanner
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9, BG_LIGHT)
    add_header(s9, "Algorithm 5: Standalone Botanical Vision Plant Scanner")

    add_card(s9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s9.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Computer Vision Feature Extraction"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    cv_pts = [
        "Lightweight, offline botanical image analyzer running on PIL and NumPy without external heavy model dependencies.",
        "Color-Space Analysis: Extracts normalized R, G, B channels from 128×128 resized leaf photographs.",
        "Excess Green Index (ExG): ExG = 2G - R - B to separate botanical leaf chlorophyll from background noise.",
        "Hue Saturation & Warmth: Calculates warmth (R-B)/(R+G+B) and HSV saturation to differentiate leaf types.",
        "Distance Metric: Standardized Euclidean distance with softmax confidence scoring."
    ]
    for pt in cv_pts:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(7)

    # Right: Supported Plants
    add_card(s9, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s9.shapes.add_textbox(Inches(7.2), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Supported Medicinal Plants & Outputs"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    plants = [
        ("Tulsi (Holy Basil)", "Ocimum tenuiflorum", "Respiratory care, tea recipe"),
        ("Neem (Indian Lilac)", "Azadirachta indica", "Blood purifier, skin rinse"),
        ("Aloe Vera (Ghritkumari)", "Aloe barbadensis", "Cooling, topical & juice"),
        ("Mint (Pudina)", "Mentha spicata", "Carminative digestive relief"),
        ("Ashwagandha", "Withania somnifera", "Rasayana, warm milk dosage"),
        ("Ginger (Adrak)", "Zingiber officinale", "Agni booster, decoction")
    ]
    for name, sci, use in plants:
        p = tf.add_paragraph()
        p.text = f"🌿 {name} ({sci}): {use}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(5)

    # =========================================================================
    # SLIDE 10: Patient-Centric UI & Safety Triage
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10, BG_LIGHT)
    add_header(s10, "Patient-Centric UI & Clinical Safety Triage")

    add_card(s10, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s10.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Emergency Safety Red-Flag Filter"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    safe_pts = [
        "Rule-based safety triage intercepts critical emergency queries immediately before recommendation.",
        "Detects severe indicators: Crushing chest pain, extreme breathlessness, severe convulsions, unconsciousness.",
        "Displays prominent warning banner advising immediate emergency medical attention (112 / 108 / 911).",
        "Suppresses home remedies during acute medical crises to ensure patient safety and ethical AI adherence."
    ]
    for pt in safe_pts:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(8)

    # Right: UI Features
    add_card(s10, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s10.shapes.add_textbox(Inches(7.2), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "User-Friendly Personalization"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    ui_pts = [
        "30-Second Dosha Discovery Quiz: 4-question interactive modal identifying Vata, Pitta, or Kapha with fallback for unknown doshas.",
        "Plain-English Severity Levels: Mild (🟢), Moderate (🟡), and Severe (🔴) with guided real-world criteria.",
        "Actionable Result Cards: Shows Condition name, % match, recommended herbs, home preparation, diet, yoga, and recovery.",
        "Zero Jargon: Completely free of confusing unit numbers or AI algorithm labels on patient-facing screens."
    ]
    for pt in ui_pts:
        p = tf.add_paragraph()
        p.text = "✔ " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 11: Testing & Evaluation Metrics
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s11, BG_LIGHT)
    add_header(s11, "Validation, Benchmarking & Test Results")

    add_card(s11, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s11.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "ML Model Benchmarking"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    benchmarks = [
        ("Random Forest", "93.50% Training Accuracy (120 trees)"),
        ("Decision Tree", "Baseline single tree comparison"),
        ("AdaBoost", "Sequential ensemble boosting"),
        ("Naive Bayes", "Multinomial text benchmark (alpha=0.5)"),
        ("Evaluation Suite", "evaluate_classification.py (80/20 train/test split)")
    ]
    for m, d in benchmarks:
        p = tf.add_paragraph()
        p.text = f"• {m}: {d}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(8)

    # Right: Test Suite
    add_card(s11, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s11.shapes.add_textbox(Inches(7.2), Inches(1.8), Inches(5.0), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Comprehensive Automated Test Suite"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    t_pts = [
        "File: test_app.py running under Python unittest.",
        "16 / 16 Integration & End-to-End Tests Passing (100%).",
        "Tests Home, Symptom Checker, Cough/Migraine queries, Emergency Red-Flag Triage, Empty queries, Herb Explorer, Condition Detail, Analytics, and Fuzzy Calculator.",
        "Runs 100% locally with zero external network dependencies."
    ]
    for pt in t_pts:
        p = tf.add_paragraph()
        p.text = "✔ " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 12: Conclusion & Future Scope
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s12, PRIMARY_DARK)

    add_card(s12, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5), bg_color=PRIMARY_DARK, border_color=LIGHT_GREEN)

    tb = s12.shapes.add_textbox(Inches(1.5), Inches(1.3), Inches(10.333), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Conclusion & Future Work"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    c_pts = [
        "Successfully developed an intelligent, data-driven Ayurvedic recommendation system unifying 5 core Data Science paradigms.",
        "Combines TF-IDF NLP similarity, Random Forest & AdaBoost ML, Bayesian uncertainty networks, Mamdani Fuzzy dosage, and Botanical Vision.",
        "Delivers a clean, intuitive, and patient-safe web interface that eliminates confusing technical jargon.",
        "Future Scope:",
        "    - Deep Learning CNN integration with Kaggle PlantVillage for 50+ herb species.",
        "    - Multilingual voice input (Speech-to-Text) for rural telemedicine access.",
        "    - Mobile Application packaging using Flutter / React Native."
    ]
    for pt in c_pts:
        p = tf.add_paragraph()
        p.text = pt if pt.startswith("    -") else "• " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = RGBColor(230, 230, 230)
        p.space_before = Pt(6)

    p_end = tf.add_paragraph()
    p_end.text = "\nThank You!  |  Questions & Discussion"
    p_end.font.size = Pt(18)
    p_end.font.bold = True
    p_end.font.color.rgb = LIGHT_GREEN
    p_end.space_before = Pt(16)

    # Save presentation
    output_path = os.path.join(os.getcwd(), "AyurHerb_Project_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation generated successfully at: {output_path}")

if __name__ == "__main__":
    create_presentation()
